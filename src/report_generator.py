"""
Report Generator — PDF and Excel export for churn prediction results.
"""
import io
import os
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.backends.backend_pdf import PdfPages
import numpy as np
import pandas as pd


# ── Colour palette ───────────────────────────────────────────
HIGH_COLOR   = "#e74c3c"
MEDIUM_COLOR = "#f39c12"
LOW_COLOR    = "#2ecc71"
BRAND_COLOR  = "#2c3e50"
ACCENT_COLOR = "#3498db"

RISK_COLORS = {"HIGH": HIGH_COLOR, "MEDIUM": MEDIUM_COLOR, "LOW": LOW_COLOR}


def _risk_color(level):
    return RISK_COLORS.get(str(level).upper(), "#95a5a6")


# ══════════════════════════════════════════════════════════════
#  PDF REPORT
# ══════════════════════════════════════════════════════════════

def generate_pdf_report(res_df: pd.DataFrame, company_name: str = "Vantive Inc.") -> bytes:
    """Return multi-page PDF as bytes."""
    buf = io.BytesIO()
    run_date = datetime.now().strftime("%B %d, %Y  %H:%M")

    # Pre-compute summary stats
    total = len(res_df)
    n_high   = int((res_df["risk_level"] == "HIGH").sum())
    n_medium = int((res_df["risk_level"] == "MEDIUM").sum())
    n_low    = int((res_df["risk_level"] == "LOW").sum())
    avg_prob = float(res_df["churn_probability"].mean()) if "churn_probability" in res_df.columns else 0.0

    with PdfPages(buf) as pdf:
        # ── Page 1: Cover / Executive Summary ──────────────────
        fig = plt.figure(figsize=(11, 8.5))

        # Header band
        ax_hdr = fig.add_axes([0, 0.88, 1, 0.12])
        ax_hdr.set_facecolor(BRAND_COLOR)
        ax_hdr.axis("off")
        ax_hdr.text(0.05, 0.55, "Customer Churn Risk Report", color="white",
                    fontsize=22, fontweight="bold", va="center", transform=ax_hdr.transAxes)
        ax_hdr.text(0.05, 0.15, f"{company_name}  ·  Generated {run_date}", color="#bdc3c7",
                    fontsize=9, va="center", transform=ax_hdr.transAxes)
        ax_hdr.text(0.97, 0.5, "CONFIDENTIAL", color="#e74c3c", fontsize=8,
                    fontweight="bold", ha="right", va="center", transform=ax_hdr.transAxes)

        # KPI boxes
        kpis = [
            ("Total Customers", f"{total:,}", ACCENT_COLOR),
            ("High Risk",        f"{n_high:,}",   HIGH_COLOR),
            ("Medium Risk",      f"{n_medium:,}", MEDIUM_COLOR),
            ("Low Risk",         f"{n_low:,}",    LOW_COLOR),
            ("Avg Churn Prob",   f"{avg_prob:.1%}", BRAND_COLOR),
        ]
        box_w, box_h = 0.16, 0.16
        for i, (label, value, color) in enumerate(kpis):
            x = 0.04 + i * (box_w + 0.025)
            y = 0.66
            ax_box = fig.add_axes([x, y, box_w, box_h])
            ax_box.set_facecolor(color)
            ax_box.set_xticks([]); ax_box.set_yticks([])
            for sp in ax_box.spines.values(): sp.set_visible(False)
            ax_box.text(0.5, 0.62, value, color="white", fontsize=18,
                        fontweight="bold", ha="center", va="center", transform=ax_box.transAxes)
            ax_box.text(0.5, 0.20, label, color="white", fontsize=8,
                        ha="center", va="center", transform=ax_box.transAxes)

        # Pie chart — risk distribution
        ax_pie = fig.add_axes([0.04, 0.12, 0.38, 0.50])
        sizes  = [n_high, n_medium, n_low]
        labels = ["High Risk", "Medium Risk", "Low Risk"]
        colors = [HIGH_COLOR, MEDIUM_COLOR, LOW_COLOR]
        non_zero = [(s, l, c) for s, l, c in zip(sizes, labels, colors) if s > 0]
        if non_zero:
            s, l, c = zip(*non_zero)
            wedges, texts, autotexts = ax_pie.pie(
                s, labels=l, colors=c, autopct="%1.1f%%",
                startangle=140, pctdistance=0.82,
                wedgeprops={"edgecolor": "white", "linewidth": 2})
            for at in autotexts: at.set_fontsize(9); at.set_color("white"); at.set_fontweight("bold")
        ax_pie.set_title("Risk Distribution", fontsize=13, fontweight="bold", pad=10)

        # Bar chart — top 10 by probability
        ax_bar = fig.add_axes([0.50, 0.12, 0.47, 0.50])
        top10 = res_df.nlargest(10, "churn_probability")
        id_col = "customer_name" if "customer_name" in top10.columns else "customer_id"
        labels_bar = top10[id_col].astype(str).str[:18].tolist()
        probs_bar  = top10["churn_probability"].tolist()
        bar_colors = [_risk_color(r) for r in top10.get("risk_level", ["HIGH"] * 10)]
        bars = ax_bar.barh(range(len(labels_bar)), probs_bar, color=bar_colors, edgecolor="white", height=0.7)
        ax_bar.set_yticks(range(len(labels_bar)))
        ax_bar.set_yticklabels(labels_bar, fontsize=8)
        ax_bar.set_xlabel("Churn Probability", fontsize=9)
        ax_bar.set_xlim(0, 1)
        ax_bar.xaxis.set_major_formatter(plt.FuncFormatter(lambda x, _: f"{x:.0%}"))
        ax_bar.set_title("Top 10 At-Risk Customers", fontsize=13, fontweight="bold")
        ax_bar.invert_yaxis()
        ax_bar.grid(axis="x", alpha=0.3)
        for bar, prob in zip(bars, probs_bar):
            ax_bar.text(min(prob + 0.01, 0.98), bar.get_y() + bar.get_height() / 2,
                        f"{prob:.1%}", va="center", fontsize=7, fontweight="bold")

        pdf.savefig(fig, bbox_inches="tight")
        plt.close(fig)

        # ── Page 2: High-Risk Customers Table ───────────────────
        high_risk = res_df[res_df["risk_level"] == "HIGH"].sort_values("churn_probability", ascending=False)
        if not high_risk.empty:
            _render_table_page(pdf, high_risk, "High-Risk Customers", HIGH_COLOR, run_date)

        # ── Page 3: Medium-Risk Customers Table ─────────────────
        medium_risk = res_df[res_df["risk_level"] == "MEDIUM"].sort_values("churn_probability", ascending=False)
        if not medium_risk.empty:
            _render_table_page(pdf, medium_risk, "Medium-Risk Customers", MEDIUM_COLOR, run_date)

        # ── Page 4: Probability Distribution Histogram ──────────
        fig, axes = plt.subplots(1, 2, figsize=(11, 5))
        _add_page_header(fig, "Churn Probability Analysis", run_date)
        plt.subplots_adjust(top=0.82, bottom=0.12, left=0.08, right=0.97, wspace=0.35)

        ax_hist = axes[0]
        probs_all = res_df["churn_probability"].dropna()
        ax_hist.hist(probs_all, bins=25, color=ACCENT_COLOR, edgecolor="white", alpha=0.85)
        ax_hist.axvline(0.4, color=MEDIUM_COLOR, linestyle="--", linewidth=1.5, label="Medium threshold (40%)")
        ax_hist.axvline(0.7, color=HIGH_COLOR,   linestyle="--", linewidth=1.5, label="High threshold (70%)")
        ax_hist.set_xlabel("Churn Probability"); ax_hist.set_ylabel("Number of Customers")
        ax_hist.set_title("Probability Distribution", fontsize=13, fontweight="bold")
        ax_hist.xaxis.set_major_formatter(plt.FuncFormatter(lambda x, _: f"{x:.0%}"))
        ax_hist.legend(fontsize=8)

        # Segment breakdown if region column exists
        ax_seg = axes[1]
        seg_col = next((c for c in ["region", "contract_type", "service_type"] if c in res_df.columns), None)
        if seg_col:
            seg = res_df.groupby(seg_col)["churn_probability"].mean().sort_values(ascending=True)
            seg.plot(kind="barh", ax=ax_seg, color=ACCENT_COLOR, edgecolor="white")
            ax_seg.set_title(f"Avg Churn Prob by {seg_col.replace('_', ' ').title()}", fontsize=13, fontweight="bold")
            ax_seg.set_xlabel("Avg Churn Probability")
            ax_seg.xaxis.set_major_formatter(plt.FuncFormatter(lambda x, _: f"{x:.0%}"))
        else:
            ax_seg.text(0.5, 0.5, "No segment column available\n(region / contract_type / service_type)",
                        ha="center", va="center", fontsize=11, color="#7f8c8d", transform=ax_seg.transAxes)
            ax_seg.axis("off")

        pdf.savefig(fig, bbox_inches="tight")
        plt.close(fig)

        # PDF metadata
        d = pdf.infodict()
        d["Title"]   = "Customer Churn Risk Report"
        d["Author"]  = company_name
        d["Subject"] = "Churn Prediction Analysis"

    buf.seek(0)
    return buf.read()


def _add_page_header(fig, title: str, run_date: str):
    ax = fig.add_axes([0, 0.92, 1, 0.06])
    ax.set_facecolor(BRAND_COLOR)
    ax.axis("off")
    ax.text(0.03, 0.5, title, color="white", fontsize=14, fontweight="bold",
            va="center", transform=ax.transAxes)
    ax.text(0.97, 0.5, run_date, color="#bdc3c7", fontsize=8,
            ha="right", va="center", transform=ax.transAxes)


def _render_table_page(pdf, df: pd.DataFrame, title: str, color: str, run_date: str):
    display_cols = [c for c in [
        "customer_id", "customer_name", "region", "account_tenure_months",
        "contract_type", "churn_probability", "risk_level", "prediction"
    ] if c in df.columns]

    rows_per_page = 28
    chunks = [df[display_cols].iloc[i:i + rows_per_page] for i in range(0, len(df), rows_per_page)]

    for page_num, chunk in enumerate(chunks):
        fig = plt.figure(figsize=(11, 8.5))
        _add_page_header(fig, f"{title}  (page {page_num + 1}/{len(chunks)})", run_date)

        ax = fig.add_axes([0.02, 0.04, 0.96, 0.84])
        ax.axis("off")

        col_labels = [c.replace("_", "\n") for c in display_cols]
        table_data = []
        for _, row in chunk.iterrows():
            formatted = []
            for c in display_cols:
                val = row[c]
                if c == "churn_probability":
                    formatted.append(f"{float(val):.1%}")
                else:
                    formatted.append(str(val) if pd.notna(val) else "—")
            table_data.append(formatted)

        tbl = ax.table(cellText=table_data, colLabels=col_labels,
                       loc="center", cellLoc="center")
        tbl.auto_set_font_size(False)
        tbl.set_fontsize(7.5)
        tbl.scale(1, 1.4)

        # Header row styling
        for j in range(len(display_cols)):
            cell = tbl[0, j]
            cell.set_facecolor(color)
            cell.set_text_props(color="white", fontweight="bold")

        # Alternate row shading + risk colouring
        risk_idx = display_cols.index("risk_level") if "risk_level" in display_cols else -1
        prob_idx = display_cols.index("churn_probability") if "churn_probability" in display_cols else -1
        for i in range(1, len(table_data) + 1):
            bg = "#f9f9f9" if i % 2 == 0 else "white"
            for j in range(len(display_cols)):
                cell = tbl[i, j]
                cell.set_facecolor(bg)
                if j == risk_idx:
                    lvl = table_data[i - 1][risk_idx]
                    cell.set_facecolor(_risk_color(lvl))
                    cell.set_text_props(color="white", fontweight="bold")
                if j == prob_idx:
                    try:
                        p = float(table_data[i - 1][prob_idx].strip("%")) / 100
                        if p >= 0.7:
                            cell.set_text_props(color=HIGH_COLOR, fontweight="bold")
                    except ValueError:
                        pass

        pdf.savefig(fig, bbox_inches="tight")
        plt.close(fig)


# ══════════════════════════════════════════════════════════════
#  EXCEL REPORT
# ══════════════════════════════════════════════════════════════

def generate_excel_report(res_df: pd.DataFrame, company_name: str = "Vantive Inc.") -> bytes:
    """Return formatted .xlsx as bytes."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import (
            PatternFill, Font, Alignment, Border, Side, numbers
        )
        from openpyxl.utils import get_column_letter
        from openpyxl.chart import BarChart, Reference, PieChart
    except ImportError:
        raise RuntimeError("openpyxl is required for Excel export. Run: pip install openpyxl")

    buf = io.BytesIO()
    wb  = Workbook()

    run_date = datetime.now().strftime("%Y-%m-%d %H:%M")
    total    = len(res_df)
    n_high   = int((res_df["risk_level"] == "HIGH").sum())
    n_medium = int((res_df["risk_level"] == "MEDIUM").sum())
    n_low    = int((res_df["risk_level"] == "LOW").sum())
    avg_prob = float(res_df["churn_probability"].mean()) if "churn_probability" in res_df.columns else 0.0

    # ── Helpers ─────────────────────────────────────────────
    def hfill(hex_str):
        return PatternFill("solid", fgColor=hex_str.lstrip("#"))

    def hfont(bold=False, color="000000", size=11):
        return Font(bold=bold, color=color.lstrip("#"), size=size)

    def center():
        return Alignment(horizontal="center", vertical="center")

    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    HEADER_FILL = hfill(BRAND_COLOR)
    HIGH_FILL   = hfill("FADBD8")
    MED_FILL    = hfill("FDEBD0")
    LOW_FILL    = hfill("D5F5E3")

    def style_header_row(ws, row, n_cols):
        for c in range(1, n_cols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill   = HEADER_FILL
            cell.font   = hfont(bold=True, color="FFFFFF")
            cell.alignment = center()
            cell.border = border

    def auto_width(ws, min_w=10, max_w=40):
        for col_cells in ws.columns:
            length = max(len(str(c.value or "")) for c in col_cells)
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, min_w), max_w)

    def row_fill_for_risk(risk):
        r = str(risk).upper()
        if r == "HIGH":   return HIGH_FILL
        if r == "MEDIUM": return MED_FILL
        return LOW_FILL

    # ── Sheet 1: Summary ────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:F1")
    title_cell = ws["A1"]
    title_cell.value     = "Customer Churn Risk Report"
    title_cell.fill      = HEADER_FILL
    title_cell.font      = hfont(bold=True, color="FFFFFF", size=16)
    title_cell.alignment = center()
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:F2")
    ws["A2"].value     = f"{company_name}  ·  {run_date}"
    ws["A2"].font      = hfont(color="7F8C8D", size=10)
    ws["A2"].alignment = center()
    ws["A2"].fill      = hfill("2C3E50")

    kpi_data = [
        ("Total Customers", total,          "3498DB"),
        ("High Risk",        n_high,        "E74C3C"),
        ("Medium Risk",      n_medium,       "F39C12"),
        ("Low Risk",         n_low,         "2ECC71"),
        ("Avg Churn Prob",   f"{avg_prob:.1%}", "8E44AD"),
    ]
    for col_i, (label, value, color) in enumerate(kpi_data, start=1):
        label_cell = ws.cell(row=4, column=col_i, value=label)
        label_cell.font      = hfont(bold=True, color="FFFFFF", size=9)
        label_cell.fill      = hfill(color)
        label_cell.alignment = center()
        label_cell.border    = border
        val_cell = ws.cell(row=5, column=col_i, value=value)
        val_cell.font        = hfont(bold=True, size=14)
        val_cell.alignment   = center()
        val_cell.border      = border
        ws.row_dimensions[5].height = 30

    # Risk breakdown mini table
    ws["A7"].value = "Risk Breakdown"; ws["A7"].font = hfont(bold=True, size=12)
    headers = ["Risk Level", "Count", "% of Total"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=8, column=ci, value=h)
        c.fill = HEADER_FILL; c.font = hfont(bold=True, color="FFFFFF"); c.alignment = center(); c.border = border

    for ri, (label, count, fill) in enumerate([
        ("High",   n_high,   HIGH_FILL),
        ("Medium", n_medium, MED_FILL),
        ("Low",    n_low,    LOW_FILL),
    ], start=9):
        ws.cell(row=ri, column=1, value=label).fill = fill
        ws.cell(row=ri, column=2, value=count).fill = fill
        pct = ws.cell(row=ri, column=3, value=count / total if total else 0)
        pct.number_format = "0.0%"; pct.fill = fill
        for ci in range(1, 4): ws.cell(row=ri, column=ci).border = border

    auto_width(ws)

    # ── Sheet 2: All Results ─────────────────────────────────
    _write_results_sheet(wb, res_df, "All Results", None, style_header_row, border, center, hfont, row_fill_for_risk)

    # ── Sheet 3: High Risk ───────────────────────────────────
    high_df = res_df[res_df["risk_level"] == "HIGH"].sort_values("churn_probability", ascending=False)
    if not high_df.empty:
        _write_results_sheet(wb, high_df, "High Risk", HIGH_FILL, style_header_row, border, center, hfont, row_fill_for_risk)

    # ── Sheet 4: Medium Risk ─────────────────────────────────
    med_df = res_df[res_df["risk_level"] == "MEDIUM"].sort_values("churn_probability", ascending=False)
    if not med_df.empty:
        _write_results_sheet(wb, med_df, "Medium Risk", MED_FILL, style_header_row, border, center, hfont, row_fill_for_risk)

    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _write_results_sheet(wb, df, sheet_name, header_fill_override, style_header_row_fn,
                          border, center_fn, hfont_fn, row_fill_fn):
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title=sheet_name)
    ws.sheet_view.showGridLines = False

    display_cols = [c for c in [
        "customer_id", "customer_name", "region", "account_tenure_months",
        "contract_type", "service_type", "churn_probability", "risk_level", "prediction"
    ] if c in df.columns]

    col_labels = {
        "customer_id":           "Customer ID",
        "customer_name":         "Customer Name",
        "region":                "Region",
        "account_tenure_months": "Tenure (Mo)",
        "contract_type":         "Contract Type",
        "service_type":          "Service Type",
        "churn_probability":     "Churn Prob",
        "risk_level":            "Risk Level",
        "prediction":            "Prediction",
    }

    for ci, col in enumerate(display_cols, 1):
        cell = ws.cell(row=1, column=ci, value=col_labels.get(col, col))
        cell.fill = header_fill_override if header_fill_override else PatternFill("solid", fgColor="2C3E50")
        from openpyxl.styles import Font, Alignment, Border
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    prob_col_idx = display_cols.index("churn_probability") + 1 if "churn_probability" in display_cols else None
    risk_col_idx = display_cols.index("risk_level") + 1      if "risk_level"          in display_cols else None

    for ri, (_, row) in enumerate(df[display_cols].iterrows(), start=2):
        risk_val = str(row.get("risk_level", "LOW")).upper()
        fill = row_fill_fn(risk_val)
        for ci, col in enumerate(display_cols, 1):
            val = row[col]
            if col == "churn_probability":
                val = round(float(val), 4)
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.fill   = fill
            cell.border = border
            if col == "churn_probability":
                cell.number_format = "0.0%"
            if col == "risk_level":
                from openpyxl.styles import Font
                cell.font = Font(bold=True)

    for col_cells in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col_cells), default=10)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(max_len + 2, 12), 38)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
